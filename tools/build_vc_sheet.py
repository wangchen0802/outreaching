"""Build collateral/SimReal-海外VC名单.xlsx from lists/investors.csv, the investor waves and the drafts.

Overseas investors only. Sheets: 可直接发邮件 (a published address), 无公开邮箱 (form, public
profile or intro), 暂不联系 (wave 9), 说明. Needs openpyxl.

Usage: python3 tools/build_vc_sheet.py
"""
import csv
import json
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).parent))
import drafts  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "collateral" / "SimReal-海外VC名单.xlsx"
EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
NONE = ("", "未找到", "无")
WAVE_LABEL = {1: "C 组：先发练手", 2: "B 组", None: "B 组", 3: "A 组：最匹配，最后发"}
WAVE_RANK = {1: 0, 2: 1, None: 1, 3: 2}
CONF_RANK = {"高": 0, "中": 1, "低": 2}
OFFICIAL_INBOX = ("pitch@", "seed@", "deals@", "submit@", "apply@", "embed@", "aistart@", "ventures@", "build@", "strategicvc@")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="111110")
BODY_FONT = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")


def blank(v):
    return (v or "").strip() in NONE


def address(r):
    """(email, type, source) for a row: the partner's own address first, then a fund inbox."""
    if not blank(r["邮箱"]):
        return r["邮箱"].strip(), "个人公开", r["邮箱来源"]
    m = EMAIL.search(r["备用渠道"] or "")
    if m:
        e = m.group(0)
        kind = "官方投递" if e.lower().startswith(OFFICIAL_INBOX) else "官方通用"
        return e, kind, r["备用渠道"]
    return "", "", ""


def how_to(r, kind):
    who = r["联系人"] if not blank(r["联系人"]) else ""
    if kind == "个人公开":
        return f"直接写给 {who}"
    if who:
        return f"第一句写明请转交 {who}（Could you pass this to {who}?）"
    return "写给团队（Hi team）"


def email_of(slug):
    p = drafts.TRACK_DIRS["investor"] / f"{slug}.md"
    if not p.exists():
        return "", ""
    email = drafts.parse(p)["parts"].get("email", "")
    subject_line, _, body = email.partition("\n")
    return subject_line.split(":", 1)[-1].strip(), body.strip()


def sheet(wb, title, cols, rows):
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = HEAD_FONT, HEAD_FILL, WRAP
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font, cell.alignment = BODY_FONT, WRAP
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions


def main():
    with open(ROOT / "lists" / "investors.csv", encoding="utf-8-sig") as f:
        rows = [r for r in csv.DictReader(f) if r["地区"].startswith("海外")]
    waves = json.loads((ROOT / "intel" / "outreach" / "investor-waves.json").read_text(encoding="utf-8"))
    for r in rows:
        w = waves.get(r["slug"], {})
        r["_wave"], r["_why"] = w.get("wave"), w.get("why", "")
        r["_email"], r["_kind"], r["_src"] = address(r)
    order = lambda r: (WAVE_RANK.get(r["_wave"], 1), CONF_RANK.get(r["置信度"], 1), r["名称"].lower())  # noqa: E731
    live = [r for r in rows if r["_wave"] != 9]
    with_email = sorted([r for r in live if r["_email"]], key=order)
    no_email = sorted([r for r in live if not r["_email"]], key=order)
    later = sorted([r for r in rows if r["_wave"] == 9], key=lambda r: r["名称"].lower())

    wb = Workbook()
    wb.remove(wb.active)
    cols = [("序号", 6), ("分组", 14), ("机构", 22), ("类型", 8), ("地区", 14), ("联系人", 18), ("职位", 26),
            ("邮箱", 30), ("邮箱类型", 10), ("写法", 30), ("开场句（英文，已写进正文）", 60), ("为什么找他们", 44),
            ("注意", 44), ("邮箱来源", 60), ("联系人来源", 44), ("审批页卡片", 22), ("邮件主题", 40), ("邮件正文", 90)]
    out = []
    for i, r in enumerate(with_email, 1):
        subject, body = email_of(r["slug"])
        out.append([i, WAVE_LABEL.get(r["_wave"], "第 2 批"), r["名称"], r["类型"], r["地区"], r["联系人"], r["职位"],
                    r["_email"], r["_kind"], how_to(r, r["_kind"]), r["钩子"], r["_why"] or r["切入点"], r["风险"],
                    r["_src"], r["联系人来源"], "v-" + r["slug"], subject, body])
    sheet(wb, "可直接发邮件", cols, out)

    cols = [("序号", 6), ("分组", 14), ("机构", 22), ("类型", 8), ("地区", 14), ("联系人", 18), ("职位", 26),
            ("备用渠道（官网表单 / 公开主页）", 60), ("开场句（英文）", 60), ("为什么找他们", 44), ("注意", 44),
            ("联系人来源", 44), ("审批页卡片", 22)]
    out = [[i, WAVE_LABEL.get(r["_wave"], "第 2 批"), r["名称"], r["类型"], r["地区"], r["联系人"], r["职位"],
            r["备用渠道"], r["钩子"], r["_why"] or r["切入点"], r["风险"], r["联系人来源"], "v-" + r["slug"]]
           for i, r in enumerate(no_email, 1)]
    sheet(wb, "无公开邮箱", cols, out)

    sheet(wb, "暂不联系", [("机构", 26), ("联系人", 22), ("原因", 80)],
          [[r["名称"], r["联系人"], r["_why"]] for r in later])

    ws = wb.create_sheet("说明")
    personal = sum(1 for r in with_email if r["_kind"] == "个人公开")
    notes = [
        f"海外投资人共 {len(rows)} 家：可直接发邮件 {len(with_email)} 家（其中合伙人本人邮箱 {personal} 家，基金官方邮箱 {len(with_email) - personal} 家），"
        f"无公开邮箱 {len(no_email)} 家，暂不联系 {len(later)} 家。",
        "分组（Luke Sophinos 冷邮件方法）：按匹配度分 C / B / A 三组，从最不匹配的 C 组先发，用回复打磨话术和 BP；B 组其次；最匹配、最想拿下的 A 组最后发。",
        "邮箱只用本人或所在基金公开发布的地址，'邮箱来源'列写了出处。没有按格式猜，没有用 RocketReach、ZoomInfo 等数据经纪网站。",
        "标'搜索摘要'的来源是在搜索结果里看到的；本环境打不开基金官网，发之前可以点开来源再确认一次。",
        "发到基金官方邮箱时，第一句写明请转交哪位合伙人（见'写法'列）。",
        "正文由 ops/templates.md 的投资人模板生成，开场句针对每家机构。跟进：第 5 天、第 12 天，模板在同一个文件。",
        "发件邮箱 business@simreal.co。审批页（https://claude.ai/artifact/WCBiAySyzEqAHkJ2YkbUj2）'投资人'标签里有同样的卡片，可以一键打开 Gmail 草稿并记录发送和跟进。",
        "无公开邮箱的机构：优先找共同联系人引荐；其次用官网表单，或在 X / LinkedIn 上手动私信本人。",
    ]
    for n in notes:
        ws.append([n])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}: {len(with_email)} with email ({personal} personal), "
          f"{len(no_email)} without, {len(later)} not now")


if __name__ == "__main__":
    main()
