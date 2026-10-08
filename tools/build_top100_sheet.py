"""Build collateral/SimReal-Top100VC-2026.xlsx and an Apollo-ready CSV from intel/investors/strebulaev-2026.csv.

The source is the 2026 Strebulaev-Jackson Venture Ranking image (top 100 US VC firms), transcribed by hand.
Rows whose logo could not be read are marked 待确认 and left out of the Apollo CSV until resolved.

Usage: python3 tools/build_top100_sheet.py
"""
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "intel" / "investors" / "strebulaev-2026.csv"
OUT = ROOT / "collateral" / "SimReal-Top100VC-2026.xlsx"
APOLLO = ROOT / "collateral" / "SimReal-Top100VC-Apollo导入.csv"


def main():
    rows = list(csv.DictReader(open(SRC, encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = "Top 100"
    cols = [("排名", 6), ("机构", 40), ("分数", 8), ("代表投资", 24), ("辨认把握", 30), ("官网域名", 24)]
    ws.append([c for c, _ in cols])
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="111110")
    for r in rows:
        ws.append([int(r["rank"]), r["firm"], int(r["score"]), r["top_deal"], r["certainty"], r["domain"]])
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font, cell.alignment = Font(name="Arial", size=10), Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    notes = wb.create_sheet("说明")
    for line in [
        "来源：2026 Strebulaev-Jackson Venture Ranking（美国前 100 家 VC），按截图逐行转录；分数和代表投资照抄图中数字。",
        "辨认把握：清楚 = logo 文字可读；较清楚 = logo 可辨但有点糊；待确认 = logo 太小或被遮挡，按代表投资推断或暂时空着。",
        "官网域名只填了有把握的，供 Apollo 匹配公司用；空着的直接按名称搜。",
        "SimReal-Top100VC-Apollo导入.csv 只含已辨认的机构（Company Name, Website 两列），可以在 Apollo 里用 CSV 导入批量查公司。",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    wb.save(OUT)
    with open(APOLLO, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Company Name", "Website"])
        for r in rows:
            if not r["certainty"].startswith("待确认"):
                w.writerow([r["firm"].split(" (")[0], r["domain"]])
    print(f"wrote {OUT.relative_to(ROOT)} and {APOLLO.relative_to(ROOT)}: "
          f"{sum(1 for r in rows if not r['certainty'].startswith('待确认'))} firms ready for Apollo")


if __name__ == "__main__":
    main()
